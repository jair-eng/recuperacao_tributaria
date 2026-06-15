from __future__ import annotations

from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.domain.relatorio_executivo.estilos_excel import style_titulo, style_header_row
from app.utils.dates import periodo_fechamento_trimestre, buscar_bloco_m_trimestre, agregar_trimestral
from app.utils.ecd_gap_utils import resolver_nat_categoria
from app.utils.excel import autosize_columns
from app.utils.numbers import to_decimal


def _pct(numerador: Decimal, denominador: Decimal) -> Decimal:
    if denominador <= 0:
        return Decimal("0.00")
    return numerador / denominador


def criar_aba_base_por_categoria(
    wb: Workbook,
    ctx: dict,
    titulo: str = "Base por Categoria - Trimestral",
) -> None:
    ws = wb.create_sheet("Base por Categoria")
    mapa_nat = ctx.get("mapa_nat_bc_cred") or {}

    agregado = agregar_trimestral(ctx)

    ws.merge_cells("A1:Q1")
    style_titulo(ws, "A1", titulo)

    headers = [
        "Ano-Trimestre",
        "Cód. Nat.",
        "Natureza",
        "Categoria",
        "Valor ECD",
        "C170 Creditado",
        "C170 Oportunidade",
        "F100 Creditado",
        "A170 Creditado",
        "Total Creditado",
        "Total Documentado",
        "GAP ECD",
        "Cobertura Documental %",
        "Qtd Contas",
        "Qtd C170",
        "Qtd F100",
        "Qtd A170",
    ]

    ws.append(headers)
    style_header_row(ws[2])

    for item in sorted(
            agregado.values(),
            key=lambda x: (x["trimestre"], x["categoria"], resolver_nat_categoria(x)),
    ):
        trimestre = item["trimestre"]
        nat = resolver_nat_categoria(item)
        categoria = item["categoria"]

        valor_ecd = to_decimal(item.get("valor_ecd"))
        valor_c170 = to_decimal(item.get("valor_creditado_c170"))
        valor_oportunidade_c170 = to_decimal(item.get("valor_oportunidade_c170"))
        valor_f100 = to_decimal(item.get("valor_creditado_f100"))
        valor_a170 = to_decimal(item.get("valor_creditado_a170"))
        valor_sem_credito_a170 = to_decimal(item.get("valor_sem_credito_a170"))

        total_creditado = valor_c170 + valor_f100 + valor_a170
        total_documentado = total_creditado + valor_sem_credito_a170
        gap_ecd = max(Decimal("0.00"), valor_ecd - total_documentado)
        cobertura = _pct(total_documentado, valor_ecd)

        ws.append([
            trimestre,
            nat,
            mapa_nat.get(nat, ""),
            categoria,
            valor_ecd,
            valor_c170,
            valor_oportunidade_c170,
            valor_f100,
            valor_a170,
            total_creditado,
            total_documentado,
            gap_ecd,
            cobertura,
            int(item.get("qtd_contas") or 0),
            int(item.get("qtd_c170") or 0),
            int(item.get("qtd_f100") or 0),
            int(item.get("qtd_a170") or 0),
        ])

    ws.auto_filter.ref = f"A2:Q{ws.max_row}"
    ws.freeze_panes = "A3"

    for row in range(3, ws.max_row + 1):
        for col in ["E", "F", "G", "H", "I", "J", "K", "L"]:
            ws[f"{col}{row}"].number_format = '#,##0.00'
        ws[f"M{row}"].number_format = '0.00%'
        for col in ["N", "O", "P", "Q"]:
            ws[f"{col}{row}"].number_format = '0'

    fill_gap = PatternFill("solid", fgColor="F4CCCC")
    for row in range(3, ws.max_row + 1):
        if to_decimal(ws[f"L{row}"].value) > 0:
            ws[f"L{row}"].fill = fill_gap

    autosize_columns(ws)

