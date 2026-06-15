from __future__ import annotations

from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from app.utils.ecd_gap_utils import pct, montar_base_por_natureza_trimestral
from app.domain.relatorio_executivo.estilos_excel import style_titulo, style_header_row
from app.utils.dates import periodo_fechamento_trimestre, buscar_bloco_m_trimestre, agregar_trimestral
from app.utils.excel import autosize_columns
from app.utils.numbers import to_decimal


def criar_aba_base_por_natureza(
    wb: Workbook,
    ctx: dict,
    titulo: str = "Base por Natureza - Trimestral",
) -> None:
    ws = wb.create_sheet("Base por Natureza")
    mapa_nat = ctx.get("mapa_nat_bc_cred") or {}

    agregado_nat = montar_base_por_natureza_trimestral(ctx)

    ws.merge_cells("A1:R1")
    style_titulo(ws, "A1", titulo)

    headers = [
        "Ano-Trimestre",
        "Cód. Nat.",
        "Natureza",
        "Valor ECD",
        "C170 Creditado",
        "C170 Oportunidade",
        "F100 Creditado",
        "A170 Creditado",
        "Total Creditado",
        "Total Documentado",
        "Bloco M Escriturado",
        "GAP ECD",
        "Dif. Documentado x M",
        "Cobertura Documental %",
        "Qtd Contas",
        "Qtd C170",
        "Qtd F100",
        "Qtd A170",
    ]

    ws.append(headers)
    style_header_row(ws[2])

    for item in sorted(
        agregado_nat.values(),
        key=lambda x: (x["trimestre"], x["nat_bc_cred"]),
    ):
        trimestre = item["trimestre"]
        nat = item["nat_bc_cred"]

        valor_ecd = to_decimal(item.get("valor_ecd"))
        valor_c170 = to_decimal(item.get("valor_creditado_c170"))
        valor_oportunidade_c170 = to_decimal(item.get("valor_oportunidade_c170"))
        valor_f100 = to_decimal(item.get("valor_creditado_f100"))
        valor_a170 = to_decimal(item.get("valor_creditado_a170"))

        total_creditado = item["total_creditado"]
        total_documentado = item["total_documentado"]
        valor_bloco_m = item["valor_bloco_m"]
        gap_ecd = item["gap_ecd"]
        dif_documentado_m = item["dif_documentado_m"]
        cobertura = item["cobertura_documental_pct"]

        ws.append([
            trimestre,
            nat,
            mapa_nat.get(nat, ""),
            valor_ecd,
            valor_c170,
            valor_oportunidade_c170,
            valor_f100,
            valor_a170,
            total_creditado,
            total_documentado,
            valor_bloco_m,
            gap_ecd,
            dif_documentado_m,
            cobertura,
            int(item.get("qtd_contas") or 0),
            int(item.get("qtd_c170") or 0),
            int(item.get("qtd_f100") or 0),
            int(item.get("qtd_a170") or 0),
        ])

    ws.auto_filter.ref = f"A2:R{ws.max_row}"
    ws.freeze_panes = "A3"

    for row in range(3, ws.max_row + 1):
        for col in ["D", "E", "F", "G", "H", "I", "J", "K", "L", "M"]:
            ws[f"{col}{row}"].number_format = '#,##0.00'
        ws[f"N{row}"].number_format = '0.00%'
        for col in ["O", "P", "Q", "R"]:
            ws[f"{col}{row}"].number_format = '0'

    fill_gap = PatternFill("solid", fgColor="F4CCCC")
    for row in range(3, ws.max_row + 1):
        if to_decimal(ws[f"L{row}"].value) > 0:
            ws[f"L{row}"].fill = fill_gap

    autosize_columns(ws)
