from __future__ import annotations

from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)

from app.utils.dates import agregar_trimestral
from app.utils.ecd_gap_utils import iter_items, ordem_categoria
from app.utils.excel import autosize_columns
from app.utils.numbers import to_decimal


def criar_aba_resumo(
    wb: Workbook,
    ctx: dict,
    *,
    titulo: str = "Revisão Fiscal PIS/COFINS - Sumário Executivo",
) -> None:

    ws = wb.create_sheet("Resumo")

    # =====================================================
    # TÍTULO
    # =====================================================

    ws.merge_cells("A1:H1")

    ws["A1"] = titulo

    ws["A1"].font = Font(
        bold=True,
        color="FFFFFF",
        size=12,
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws["A1"].fill = PatternFill(
        "solid",
        fgColor="000080",
    )

    # =====================================================
    # CABEÇALHO PRINCIPAL
    # =====================================================

    headers = [
        "Categoria",
        "Despesa Contábil Total",
        "EFD Documentada",
        "Gap Identificado",
        "Base Recuperável C170",
        "PIS Recuperável C170",
        "COFINS Recuperável C170",
        "Crédito Recuperável C170",
    ]

    ws.append(headers)

    ws.freeze_panes = "A3"

    header_fill = PatternFill(
        "solid",
        fgColor="D9D9D9",
    )

    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    # =====================================================
    # DESPESA CONTÁBIL POR CATEGORIA
    # =====================================================

    por_categoria = defaultdict(
        lambda: {
            "despesa": Decimal("0.00"),
        }
    )

    for item in iter_items(
        ctx.get("linhas_ecd")
    ):

        categoria = (
            item.get("categoria")
            or "NaoClassificado"
        )

        valor = to_decimal(
            item.get("valor")
        )

        if valor <= 0:
            continue

        por_categoria[
            categoria
        ]["despesa"] += valor

    # =====================================================
    # EFD DOCUMENTADA POR CATEGORIA
    # =====================================================

    documentado_por_categoria = defaultdict(
        lambda: Decimal("0.00")
    )

    agregado = agregar_trimestral(ctx)

    for item in agregado.values():

        categoria = (
            item.get("categoria")
            or "NaoClassificado"
        )

        total_creditado = (
            to_decimal(
                item.get(
                    "valor_creditado_c170"
                )
            )
            + to_decimal(
                item.get(
                    "valor_creditado_f100"
                )
            )
            + to_decimal(
                item.get(
                    "valor_creditado_a170"
                )
            )
        )

        total_documentado = (
            total_creditado
            + to_decimal(
                item.get(
                    "valor_oportunidade_c170"
                )
            )
            + to_decimal(
                item.get(
                    "valor_sem_credito_a170"
                )
            )
        )

        documentado_por_categoria[
            categoria
        ] += total_documentado

    # =====================================================
    # OPORTUNIDADES C170
    # =====================================================

    oportunidades_por_categoria = defaultdict(
        lambda: {
            "base_c170": Decimal("0.00"),
            "pis_c170": Decimal("0.00"),
            "cofins_c170": Decimal("0.00"),
            "credito_c170": Decimal("0.00"),
        }
    )

    for op in (
        ctx.get("oportunidades_c170")
        or []
    ):

        categoria = (
            op.get("categoria")
            or "NaoClassificado"
        )

        oportunidades_por_categoria[
            categoria
        ]["base_c170"] += to_decimal(
            op.get("base_recuperavel")
        )

        oportunidades_por_categoria[
            categoria
        ]["pis_c170"] += to_decimal(
            op.get("pis_recuperavel")
        )

        oportunidades_por_categoria[
            categoria
        ]["cofins_c170"] += to_decimal(
            op.get("cofins_recuperavel")
        )

        oportunidades_por_categoria[
            categoria
        ]["credito_c170"] += to_decimal(
            op.get("credito_recuperavel")
        )

    # =====================================================
    # DEBUG
    # =====================================================

    print(
        "[CATEGORIAS SEM ECD]",
        sorted(
            set(
                oportunidades_por_categoria.keys()
            )
            - set(
                por_categoria.keys()
            )
        ),
    )

    # =====================================================
    # TOTAIS C170
    # =====================================================

    total_ecd = Decimal("0.00")
    total_documentado = Decimal("0.00")
    total_gap = Decimal("0.00")

    total_base_c170 = Decimal("0.00")
    total_pis_c170 = Decimal("0.00")
    total_cofins_c170 = Decimal("0.00")
    total_credito_c170 = Decimal("0.00")

    todas_categorias = (
        set(
            por_categoria.keys()
        )
        | set(
            documentado_por_categoria.keys()
        )
        | set(
            oportunidades_por_categoria.keys()
        )
    )

    # =====================================================
    # LINHAS POR CATEGORIA
    # =====================================================

    for categoria in sorted(
        todas_categorias,
        key=lambda c: (
            c == "NaoClassificado",
            -por_categoria.get(
                c,
                {},
            ).get(
                "despesa",
                Decimal("0.00"),
            ),
            c,
        ),
    ):

        dados = (
            por_categoria.get(categoria)
            or {
                "despesa": Decimal("0.00")
            }
        )

        despesa = dados["despesa"]

        efd_documentado = (
            documentado_por_categoria.get(
                categoria
            )
            or Decimal("0.00")
        )

        gap = max(
            Decimal("0.00"),
            despesa - efd_documentado,
        )

        c170 = (
            oportunidades_por_categoria.get(
                categoria
            )
            or {}
        )

        base_c170 = (
            c170.get("base_c170")
            or Decimal("0.00")
        )

        pis_c170 = (
            c170.get("pis_c170")
            or Decimal("0.00")
        )

        cofins_c170 = (
            c170.get("cofins_c170")
            or Decimal("0.00")
        )

        credito_c170 = (
            c170.get("credito_c170")
            or Decimal("0.00")
        )

        total_ecd += despesa

        total_documentado += (
            efd_documentado
        )

        total_gap += gap

        total_base_c170 += (
            base_c170
        )

        total_pis_c170 += (
            pis_c170
        )

        total_cofins_c170 += (
            cofins_c170
        )

        total_credito_c170 += (
            credito_c170
        )

        ws.append([
            categoria,
            despesa,
            (
                efd_documentado
                if efd_documentado
                else ""
            ),
            (
                gap
                if gap
                else ""
            ),
            (
                base_c170
                if base_c170
                else ""
            ),
            (
                pis_c170
                if pis_c170
                else ""
            ),
            (
                cofins_c170
                if cofins_c170
                else ""
            ),
            (
                credito_c170
                if credito_c170
                else ""
            ),
        ])

    # =====================================================
    # TOTAL GERAL DA TABELA C170
    # =====================================================

    total_row = (
        ws.max_row + 1
    )

    ws.append([
        "TOTAL GERAL",
        total_ecd,
        total_documentado,
        total_gap,
        total_base_c170,
        total_pis_c170,
        total_cofins_c170,
        total_credito_c170,
    ])

    for cell in ws[total_row]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="FFF2CC",
        )

    thin = Side(
        style="thin",
        color="000000",
    )

    for cell in ws[total_row]:

        cell.border = Border(
            top=thin
        )

    # =====================================================
    # F100 - FRETES NÃO ESCRITURADOS
    # =====================================================

    resumo_f100 = (
        ctx.get(
            "resumo_f100_recuperaveis"
        )
        or {}
    )

    credito_f100 = to_decimal(
        resumo_f100.get(
            "credito_total"
        )
    )

    quantidade_f100 = int(
        resumo_f100.get(
            "quantidade"
        )
        or 0
    )

    # =====================================================
    # CRÉDITO POTENCIAL TOTAL
    # =====================================================

    credito_potencial_total = (
        total_credito_c170
        + credito_f100
    )

    # =====================================================
    # QUADRO EXECUTIVO DE CRÉDITOS
    # =====================================================

    ws.append([])

    linha_titulo_creditos = (
        ws.max_row + 1
    )

    ws.append([
        "Resumo dos Créditos Recuperáveis"
    ])

    ws.merge_cells(
        start_row=linha_titulo_creditos,
        start_column=1,
        end_row=linha_titulo_creditos,
        end_column=3,
    )

    celula_titulo = ws.cell(
        row=linha_titulo_creditos,
        column=1,
    )

    celula_titulo.font = Font(
        bold=True,
        color="FFFFFF",
    )

    celula_titulo.fill = PatternFill(
        "solid",
        fgColor="4472C4",
    )

    celula_titulo.alignment = Alignment(
        horizontal="left",
    )

    # -----------------------------------------------------
    # C170
    # -----------------------------------------------------

    ws.append([
        "Crédito Recuperável C170",
        total_credito_c170,
        "Itens/documentos fiscais analisados",
    ])

    linha_credito_c170 = ws.max_row

    # -----------------------------------------------------
    # F100
    # -----------------------------------------------------

    ws.append([
        "Crédito Recuperável F100 - Fretes",
        credito_f100,
        (
            f"{quantidade_f100} contrato(s) "
            f"não escriturado(s)"
            if quantidade_f100
            else "Nenhum contrato não escriturado"
        ),
    ])

    linha_credito_f100 = ws.max_row

    # -----------------------------------------------------
    # TOTAL POTENCIAL
    # -----------------------------------------------------

    ws.append([
        "CRÉDITO POTENCIAL TOTAL",
        credito_potencial_total,
        "C170 + F100",
    ])

    linha_credito_total = (
        ws.max_row
    )

    for cell in ws[
        linha_credito_total
    ]:

        cell.font = Font(
            bold=True,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="E2F0D9",
        )

    # =====================================================
    # FORMATA VALORES DO QUADRO
    # =====================================================

    for linha in [
        linha_credito_c170,
        linha_credito_f100,
        linha_credito_total,
    ]:

        ws.cell(
            row=linha,
            column=2,
        ).number_format = '#,##0.00'

    # =====================================================
    # NOTAS METODOLÓGICAS
    # =====================================================

    ws.append([])
    ws.append([
        "Notas Metodológicas"
    ])

    ws[
        f"A{ws.max_row}"
    ].font = Font(
        bold=True
    )

    ws.append([
        "1. Despesa Contábil Total representa os valores identificados na ECD por categoria fiscal."
    ])

    ws.append([
        "2. EFD Documentada representa os valores identificados na EFD Contribuições vinculados às operações analisadas."
    ])

    ws.append([
        "3. Gap Identificado representa diferença entre despesa elegível e valor documentado na EFD, sujeito à validação fiscal e documental."
    ])

    ws.append([
        "4. Valores recuperáveis C170 representam oportunidades documentadas por item, nota fiscal e cenário fiscal."
    ])

    ws.append([
        "5. Crédito Recuperável C170 não é estimativa teórica; decorre da aba 'Oportunidades C170'."
    ])

    ws.append([
        "6. Crédito Recuperável F100 representa o potencial identificado em contratos de frete não localizados na EFD-Contribuições."
    ])

    ws.append([
        "7. O Crédito Potencial Total corresponde à soma das oportunidades C170 e F100 identificadas no processamento."
    ])

    ws.append([
        "8. Os valores apresentados estão sujeitos à validação documental e fiscal antes de eventual aproveitamento."
    ])

    ws.append([
        "9. Itens sem lastro suficiente devem ser revisados nas abas específicas de investigação."
    ])

    # =====================================================
    # ALERTAS DE OMISSÃO
    # =====================================================

    alertas_omissao = (
        ctx.get(
            "alertas_efd_omissao"
        )
        or []
    )

    meses = sorted({
        str(
            a.get("Período")
        )
        for a in alertas_omissao
        if a.get("Período")
    })

    gap_omissao = sum(
        (
            to_decimal(
                a.get("GAP")
            )
            for a in alertas_omissao
        ),
        Decimal("0.00"),
    )

    if (
        meses
        and gap_omissao > 0
    ):

        texto_alerta = (
            f"⚠ ALERTA - "
            f"{len(meses)} mês(es) "
            f"com EFD entregue zerada "
            f"e despesa elegível na ECD: "
            f"{', '.join(meses)}. "
            f"Gap identificado: "
            f"R$ {gap_omissao:,.2f}. "
            f"Ver aba 'Alertas EFD Omissão'."
        )

        alerta_row = (
            ws.max_row + 2
        )

        ws.cell(
            row=alerta_row,
            column=1,
            value=texto_alerta,
        )

        ws.merge_cells(
            start_row=alerta_row,
            start_column=1,
            end_row=alerta_row,
            end_column=8,
        )

        cell = ws.cell(
            row=alerta_row,
            column=1,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="C00000",
        )

    # =====================================================
    # FORMATAÇÃO NUMÉRICA DA TABELA PRINCIPAL
    # =====================================================

    for col in [
        "B",
        "C",
        "D",
        "E",
        "F",
        "G",
        "H",
    ]:

        for row in range(
            3,
            total_row + 1,
        ):

            ws[
                f"{col}{row}"
            ].number_format = (
                '#,##0.00'
            )

    # =====================================================
    # LARGURAS
    # =====================================================

    autosize_columns(ws)

    larguras = {
        "A": 34,
        "B": 22,
        "C": 22,
        "D": 20,
        "E": 22,
        "F": 22,
        "G": 24,
        "H": 24,
    }

    for col, width in (
        larguras.items()
    ):

        ws.column_dimensions[
            col
        ].width = width