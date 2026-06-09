from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from decimal import Decimal
from app.domain.sped.contextos.dominios.transp import OPORTUNIDADES_TRANSP, avaliar_oportunidades_transp
from app.utils.excel import autosize_columns
from app.utils.numbers import to_decimal



def criar_aba_diagnostico_transp(wb: Workbook, ctx: dict) -> None:
    ws = wb.create_sheet("Diagnóstico TRANSP")

    ws.merge_cells("A1:H1")
    ws["A1"] = "Diagnóstico por Domínio - Transportadora"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A1"].fill = PatternFill("solid", fgColor="000080")

    ws.append([
        "Oportunidade",
        "Fontes esperadas",
        "Cobertura do Catálogo",
        "Indícios de recuperação",
        "Restrições/validações",
        "Próxima ação",
        "Situação",
        "Sinal encontrado na ECD",
    ])

    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")


    catalogo = ctx.get("catalogo_fiscal")
    linhas_ecd = ctx.get("linhas_ecd") or []

    for op in avaliar_oportunidades_transp(ctx):
        slugs = op.get("slugs") or []
        itens_catalogo = []

        grupos_fiscais = []

        for slug in slugs:
            codigos = catalogo.codigos(slug) if catalogo else []

            grupos_fiscais.append(
                f"{slug} ({len(codigos)})"
            )
            itens_catalogo.extend([str(c) for c in codigos if c])

        texto_grupos = "\n".join(grupos_fiscais)

        categorias_ecd = op.get("categorias_ecd") or []
        contas_encontradas = op.get("contas_ecd") or []
        valor_total = op.get("valor_ecd") or Decimal("0.00")

        if contas_encontradas:
            sinais = (
                    f"{op.get('qtd_contas_ecd')} conta(s) encontrada(s)\n"
                    f"Valor total: R$ {valor_total:,.2f}\n\n"
                    + "\n".join(
                f"{c.get('cod_cta')} - {c.get('nome_cta')}"
                for c in contas_encontradas[:5]
            )
            )
        else:
            sinais = "Não identificado na ECD"

        extras = op.get("sinais_extra") or []

        if extras:
            sinais = sinais + "\n\n" + "\n\n".join(extras)

        ws.append([
            op["oportunidade"],
            op["fontes"],
            op["cobertura_catalogo"],
            op["indicadores"],
            op["restricoes"],
            op["acao"],
            op["situacao"],
            sinais,
        ])

    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    autosize_columns(ws)


