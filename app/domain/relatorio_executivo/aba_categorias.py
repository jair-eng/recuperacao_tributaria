from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.config.settings import ALIQUOTA_PIS, ALIQUOTA_COFINS
from app.utils.excel import criar_aba_generica
from app.utils.numbers import to_decimal
import re
from openpyxl.styles import Font, PatternFill, Alignment


def _nome_aba_categoria(categoria: str) -> str:
    nome = str(categoria or "").strip()
    nome = nome.replace("/", "-").replace("\\", "-")
    return f"CAT_{nome[:25]}"

def categoria_legivel(categoria: str) -> str:
    if not categoria:
        return ""

    texto = re.sub(r"([a-z])([A-Z])", r"\1 \2", categoria)

    ajustes = {
        "PJ": "PJ",
        "PF": "PF",
        "MOPP": "MOPP",
    }

    partes = []
    for p in texto.split():
        partes.append(ajustes.get(p, p))

    return " ".join(partes)

def criar_abas_por_categoria(
    wb: Workbook,
    ctx: dict,
) -> None:
    por_categoria = defaultdict(list)

    # Base nova: período + nat + categoria
    por_periodo_chave = ctx.get("por_periodo_chave") or {}

    # Total ECD por período/categoria para rateio proporcional
    total_ecd_periodo_categoria = defaultdict(Decimal)
    doc_periodo_categoria = defaultdict(Decimal)

    for item in por_periodo_chave.values():
        periodo = item.get("periodo")
        categoria = item.get("categoria") or "NaoClassificado"

        if not periodo or categoria == "NaoClassificado":
            continue

        chave_pc = (periodo, categoria)

        total_ecd_periodo_categoria[chave_pc] += to_decimal(item.get("valor_ecd"))

        doc_periodo_categoria[chave_pc] += to_decimal(
            item.get("valor_documentado_total")
        )

    for item in ctx.get("linhas_ecd", []):
        categoria = item.get("categoria")

        if not categoria or categoria == "NaoClassificado":
            continue

        if item.get("fundamento") == "Investigar":
            continue

        por_categoria[categoria].append(item)

    for categoria, itens in por_categoria.items():
        rows = []

        for item in itens:
            periodo = item.get("periodo")
            valor = to_decimal(item.get("valor"))

            chave_pc = (periodo, categoria)

            total_ecd_categoria_periodo = total_ecd_periodo_categoria.get(
                chave_pc,
                Decimal("0.00"),
            )

            total_documentado_categoria_periodo = doc_periodo_categoria.get(
                chave_pc,
                Decimal("0.00"),
            )

            if total_ecd_categoria_periodo > 0:
                base_documentada = (
                    valor / total_ecd_categoria_periodo
                ) * total_documentado_categoria_periodo
            else:
                base_documentada = Decimal("0.00")

            base_documentada = base_documentada.quantize(Decimal("0.01"))

            gap = max(
                Decimal("0.00"),
                valor - base_documentada,
            ).quantize(Decimal("0.01"))

            rows.append(
                {
                    "Período": periodo,
                    "Código Conta": item.get("cod_cta"),
                    "Descrição": item.get("nome_cta"),
                    "Grupo": item.get("grupo"),
                    "Fundamento": item.get("fundamento"),
                    "Naturezas Esperadas": item.get("naturezas_esperadas"),
                    "Despesa Contábil": valor,
                    "Base Documentada(Rateada)": base_documentada,
                    "Gap ECD x Documentação": gap,
                    "Fonte": item.get("origem_classificacao") or "Heurística",
                    "Confiança": item.get("confianca"),
                    "Observação": item.get("observacao"),
                }
            )

        criar_aba_generica(
            wb,
            nome_aba=_nome_aba_categoria(categoria),
            headers=[
                "Período",
                "Código Conta",
                "Descrição",
                "Grupo",
                "Fundamento",
                "Naturezas Esperadas",
                "Despesa Contábil",
                "Base Documentada",
                "Gap ECD x Documentação",
                "Fonte",
                "Confiança",
                "Observação",
            ],
            rows=rows,
            money_cols=[
                "Despesa Contábil",
                "Base Documentada",
                "Gap ECD x Documentação",
            ],
        )

        ws = wb[_nome_aba_categoria(categoria)]

        ws.insert_rows(1, amount=2)

        ultima_coluna = ws.max_column

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=ultima_coluna,
        )

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=ultima_coluna,
        )

        ws["A1"] = f"Categoria: {categoria_legivel(categoria)}"

        ws["A2"] = (
            f"Fundamento: {itens[0].get('fundamento')} | "
            f"Naturezas Esperadas: "
            f"{', '.join(map(str, itens[0].get('naturezas_esperadas') or []))}"
        )

        ultima_coluna = ws.max_column
        ultima_linha = ws.max_row

        ws.auto_filter.ref = f"A3:{get_column_letter(ultima_coluna)}{ultima_linha}"
        ws.freeze_panes = "A4"

        fill = PatternFill(
            fill_type="solid",
            fgColor="E7E6E6",
        )

        font = Font(
            color="000000",
            bold=True,
        )

        align = Alignment(
            horizontal="center",
            vertical="center",
        )

        for celula in ("A1", "A2"):
            ws[celula].fill = fill
            ws[celula].font = font
            ws[celula].alignment = align