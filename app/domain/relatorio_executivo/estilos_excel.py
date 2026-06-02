from __future__ import annotations

from decimal import Decimal
from typing import Any, Dict, List
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BRL_NUMBER_FORMAT = '[$R$-pt-BR] #,##0.00'

def style_header(ws, row: int, col_start: int, col_end: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border

    ws.row_dimensions[row].height = 22

def style_titulo(ws, cell_ref: str, texto: str, fill: str = "000080") -> None:
    ws[cell_ref] = texto
    ws[cell_ref].font = Font(bold=True, color="FFFFFF", size=12)
    ws[cell_ref].alignment = Alignment(horizontal="center")
    ws[cell_ref].fill = PatternFill("solid", fgColor=fill)


def style_header_row(row) -> None:
    fill = PatternFill("solid", fgColor="D9D9D9")
    for cell in row:
        cell.font = Font(bold=True)
        cell.fill = fill


def apply_table_style(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    zebra = PatternFill("solid", fgColor="F7F7F7")

    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if r % 2 == 0:
                cell.fill = zebra

def _excel_value(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return str(value)

    return value


def write_rows(ws, rows: List[Dict[str, Any]], columns: List[str], money_cols: List[str]) -> None:
    ws.append(columns)
    style_header(ws, 1, 1, len(columns))

    money_idx = {columns.index(col) + 1 for col in money_cols if col in columns}

    for row in rows:
        ws.append([_excel_value(row.get(col, "")) for col in columns])

    for r in range(2, ws.max_row + 1):
        for c in money_idx:
            ws.cell(r, c).number_format = BRL_NUMBER_FORMAT

    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"
    apply_table_style(ws, 1, ws.max_row, 1, len(columns))
    autosize_columns(ws, max_width=80)

def autosize_columns(ws, max_width: int = 60) -> None:
    dims: Dict[int, int] = {}
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if cell.value is None:
                continue
            dims[cell.column] = max(dims.get(cell.column, 0), len(str(cell.value)))

    for col, w in dims.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, w + 2), max_width)


def fmt_competencia(comp: str) -> str:
    comp = str(comp or "").strip()
    if len(comp) == 6 and comp.isdigit():
        return f"01-{comp[4:6]}-{comp[:4]}"
    return comp or "000000"


