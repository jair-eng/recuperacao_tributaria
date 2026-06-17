from __future__ import annotations

from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.utils.numbers import to_decimal


def criar_aba_oportunidades_c170(
    wb: Workbook,
    ctx: dict,
) -> None:
    oportunidades = ctx.get("oportunidades_c170") or []

    credito_escriturado = sum(
        to_decimal(x.get("credito_recuperavel"))
        for x in oportunidades
        if x.get("status_cruzamento") != "NAO_ESCRITURADO"
    )

    credito_nao_declarado = sum(
        to_decimal(x.get("credito_recuperavel"))
        for x in oportunidades
        if x.get("status_cruzamento") == "NAO_ESCRITURADO"
    )

    if "Oportunidades C170" in wb.sheetnames:
        del wb["Oportunidades C170"]

    ws = wb.create_sheet("Oportunidades C170")

    moeda_cols = {
        "Base Recuperável",
        "PIS Recuperável",
        "COFINS Recuperável",
        "Crédito Recuperável",
    }

    colunas = [
        "Período",
        "Tipo Oportunidade",
        "Diagnóstico",
        "Categoria",
        "Natureza",
        "Cód. Crédito",
        "CFOP",
        "CST PIS Atual",
        "CST PIS Corrigido",
        "CST COFINS Atual",
        "CST COFINS Corrigido",
        "Cod Item",
        "Descrição",
        "NCM",
        "Cod Conta",
        "Base Recuperável",
        "PIS Recuperável",
        "COFINS Recuperável",
        "Crédito Recuperável",
        "Motivo",
        "Chave NF-e",
        "Nº Doc",
        "Nº Item",
    ]

    total_qtd = len(oportunidades)
    total_base = sum((i.get("base_recuperavel") or Decimal("0")) for i in oportunidades)
    total_pis = sum((i.get("pis_recuperavel") or Decimal("0")) for i in oportunidades)
    total_cofins = sum((i.get("cofins_recuperavel") or Decimal("0")) for i in oportunidades)
    total_credito = sum((i.get("credito_recuperavel") or Decimal("0")) for i in oportunidades)

    # =========================
    # Estilos
    # =========================
    fill_titulo = PatternFill("solid", fgColor="1F4E78")
    fill_header = PatternFill("solid", fgColor="D9EAF7")
    fill_resumo = PatternFill("solid", fgColor="E2F0D9")

    font_titulo = Font(color="FFFFFF", bold=True, size=13)
    font_header = Font(bold=True)
    font_resumo = Font(bold=True)

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # =========================
    # Título
    # =========================
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(colunas))
    cell = ws.cell(row=1, column=1, value="Oportunidades C170 ICMS x EFD Contribuições")
    cell.fill = fill_titulo
    cell.font = font_titulo
    cell.alignment = Alignment(horizontal="center")

    # =========================
    # Resumo no topo
    # =========================
    resumo = [
        ("Total de Oportunidades", total_qtd),
        ("Base Recuperável", total_base),
        ("À Corrigir", credito_escriturado ),
        ("Não Escriturado", credito_nao_declarado ),
        ("PIS Recuperável", total_pis),
        ("COFINS Recuperável", total_cofins),
        ("Crédito Recuperável", total_credito),
    ]

    for idx, (label, valor) in enumerate(resumo, start=3):
        ws.cell(row=idx, column=1, value=label)
        ws.cell(row=idx, column=2, value=valor)

        for col in (1, 2):
            c = ws.cell(row=idx, column=col)
            c.fill = fill_resumo
            c.font = font_resumo
            c.border = border

        if label != "Total de Oportunidades":
            ws.cell(row=idx, column=2).number_format = '#,##0.00'

    # =========================
    # Cabeçalho da tabela
    # =========================
    header_row = 11

    for col_idx, nome_coluna in enumerate(colunas, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=nome_coluna)
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")


    # =========================
    # Linhas
    # =========================
    for row_idx, item in enumerate(oportunidades, start=header_row + 1):


        row = [
            item.get("periodo"),
            item.get("status_oportunidade"),
            item.get("codigo_diagnostico"),
            item.get("categoria"),
            item.get("nat_bc_cred"),
            item.get("cod_cred"),
            item.get("cfop"),

            item.get("cst_pis_atual"),
            item.get("cst_pis_destino"),
            item.get("cst_cofins_atual"),
            item.get("cst_cofins_destino"),

            item.get("cod_item"),
            item.get("descricao"),
            item.get("ncm"),
            item.get("cod_cta"),

            item.get("base_recuperavel"),
            item.get("pis_recuperavel"),
            item.get("cofins_recuperavel"),
            item.get("credito_recuperavel"),

            item.get("motivo"),

            item.get("chv_nfe"),
            item.get("num_doc"),
            item.get("num_item"),
        ]

        for col_idx, valor in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = border
            cell.alignment = Alignment(vertical="top")

            nome_coluna = colunas[col_idx - 1]
            if nome_coluna in moeda_cols:
                cell.number_format = '#,##0.00'

    # =========================
    # Filtro / congelamento
    # =========================
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(colunas))}"
        f"{max(header_row + 1, header_row + total_qtd)}"
    )
    ws.freeze_panes = f"A{header_row + 1}"

    # =========================
    # Larguras
    # =========================
    larguras = {
        "A": 19,  # Período
        "B": 19,  # Tipo Oportunidade

        "C": 32,  # Diagnóstico
        "D": 28,  # Categoria

        "E": 7,  # Natureza
        "F": 8,  # Cod Crédito
        "G": 8,  # CFOP
        "H": 8,  # CST PIS Atual

        "I": 8,  # CST PIS Corrigido
        "J": 8,  # CST COFINS Atual
        "K": 8,  # CST COFINS Corrigido
        "L": 14,  # Cod Item

        "M": 45,  # Descrição
        "N": 10,  # NCM
        "O": 8,  # Cod Conta
        "P": 14,  # Base Recuperável

        "Q": 14,  # PIS Recuperável
        "R": 14,  # COFINS Recuperável
        "S": 16,  # Crédito Recuperável
        "T": 60,  # Motivo

        "U": 46,  # Chave NF-e

        "V": 10,  # Nº Doc
        "W": 8,  # Nº Item

    }

    for col, width in larguras.items():
        ws.column_dimensions[col].width = width