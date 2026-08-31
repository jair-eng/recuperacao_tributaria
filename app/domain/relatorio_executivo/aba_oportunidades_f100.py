from __future__ import annotations

from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def criar_aba_oportunidades_f100(
    wb: Workbook,
    ctx: dict[str, Any],
) -> None:

    registros = (
        ctx.get("f100_recuperaveis")
        or []
    )

    nome_aba = "F100_RECUPERAVEIS"

    if nome_aba in wb.sheetnames:
        del wb[nome_aba]

    ws = wb.create_sheet(
        title=nome_aba
    )

    # =====================================================
    # CABEÇALHO
    # =====================================================

    headers = [
        "COMPETÊNCIA",
        "NÚMERO F100",
        "DATA",
        "CONTRATANTE",
        "CNPJ CONTRATANTE",
        "CONTRATADO",
        "CPF/CNPJ CONTRATADO",
        "TIPO PESSOA",
        "VALOR FRETE",
        "CST PIS",
        "ALÍQ. PIS",
        "CRÉDITO PIS",
        "CST COFINS",
        "ALÍQ. COFINS",
        "CRÉDITO COFINS",
        "NAT BC CRED",
        "COD CRED",
        "CRÉDITO TOTAL",
        "ORIGEM",
    ]

    ws.append(headers)

    # =====================================================
    # ESTILO DO CABEÇALHO
    # =====================================================

    for celula in ws[1]:

        celula.font = Font(
            bold=True,
            color="FFFFFF",
        )

        celula.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # =====================================================
    # DADOS
    # =====================================================

    for item in registros:

        ws.append([
            item.get("competencia"),
            item.get("numero_f100"),
            item.get("data"),

            item.get("contratante"),
            item.get("documento_contratante"),

            item.get("contratado"),
            item.get("documento_contratado"),
            item.get("tipo_pessoa"),

            item.get("valor_frete"),

            item.get("cst_pis"),
            item.get("aliq_pis"),
            item.get("credito_pis"),

            item.get("cst_cofins"),
            item.get("aliq_cofins"),
            item.get("credito_cofins"),

            item.get("nat_bc_cred"),
            item.get("cod_cred"),

            item.get("credito_total"),

            item.get("origem"),
        ])

    # =====================================================
    # FORMATAÇÃO NUMÉRICA
    # =====================================================

    colunas_monetarias = [
        9,   # VALOR FRETE
        12,  # CRÉDITO PIS
        15,  # CRÉDITO COFINS
        18,  # CRÉDITO TOTAL
    ]

    for coluna in colunas_monetarias:

        for linha in range(
            2,
            ws.max_row + 1,
        ):

            ws.cell(
                row=linha,
                column=coluna,
            ).number_format = '#,##0.00'

    # =====================================================
    # FORMATAÇÃO DAS ALÍQUOTAS
    # =====================================================

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        ws.cell(
            row=linha,
            column=11,
        ).number_format = '0.0000'

        ws.cell(
            row=linha,
            column=14,
        ).number_format = '0.0000'

    # =====================================================
    # FILTRO E CONGELAMENTO
    # =====================================================

    if ws.max_row >= 1:

        ws.auto_filter.ref = (
            f"A1:S{ws.max_row}"
        )

    ws.freeze_panes = "A2"

    # =====================================================
    # LARGURA DAS COLUNAS
    # =====================================================

    larguras = {
        1: 12,
        2: 14,
        3: 12,
        4: 35,
        5: 20,
        6: 35,
        7: 20,
        8: 12,
        9: 16,
        10: 12,
        11: 12,
        12: 16,
        13: 14,
        14: 14,
        15: 18,
        16: 14,
        17: 12,
        18: 18,
        19: 35,
    }

    for coluna, largura in larguras.items():

        letra = get_column_letter(
            coluna
        )

        ws.column_dimensions[
            letra
        ].width = largura

    # =====================================================
    # TOTAL
    # =====================================================

    if registros:

        linha_total = (
            ws.max_row + 2
        )

        ws.cell(
            row=linha_total,
            column=8,
            value="TOTAL",
        )

        ws.cell(
            row=linha_total,
            column=8,
        ).font = Font(
            bold=True
        )

        resumo = (
            ctx.get(
                "resumo_f100_recuperaveis"
            )
            or {}
        )

        ws.cell(
            row=linha_total,
            column=9,
            value=resumo.get(
                "base",
                Decimal("0.00"),
            ),
        )

        ws.cell(
            row=linha_total,
            column=12,
            value=resumo.get(
                "pis",
                Decimal("0.00"),
            ),
        )

        ws.cell(
            row=linha_total,
            column=15,
            value=resumo.get(
                "cofins",
                Decimal("0.00"),
            ),
        )

        ws.cell(
            row=linha_total,
            column=18,
            value=resumo.get(
                "credito_total",
                Decimal("0.00"),
            ),
        )

        for coluna in [
            9,
            12,
            15,
            18,
        ]:

            celula = ws.cell(
                row=linha_total,
                column=coluna,
            )

            celula.number_format = (
                '#,##0.00'
            )

            celula.font = Font(
                bold=True
            )