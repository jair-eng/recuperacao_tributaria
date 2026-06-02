from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.domain.relatorio_executivo.estilos_excel import style_titulo, style_header_row
from app.utils.ecd_gap_status_utils import iter_items
from app.utils.excel import autosize_columns





def criar_aba_bases_efd_por_natureza(
    wb: Workbook,
    ctx: dict,
    titulo: str = "Bases declaradas na EFD-Contribuições por mês e natureza",
) -> None:
    ws = wb.create_sheet("BasesEfdPorNatureza")
    mapa_nat = ctx.get("mapa_nat_bc_cred") or {}

    # BLOCO 1
    ws.merge_cells("A1:D1")
    style_titulo(ws, "A1", titulo)

    headers1 = ["Ano-Mês", "Cód.", "Natureza", "Base Declarada"]
    ws.append(headers1)
    style_header_row(ws[2])

    efd_por_nat = {}

    for item in iter_items(ctx.get("por_natureza")):
        nat = str(item.get("nat_bc_cred") or "").zfill(2)
        valor_efd = item.get("valor_efd") or 0
        efd_por_nat[nat] = valor_efd

        ws.append([
            item.get("periodo"),
            nat,
            mapa_nat.get(nat, ""),
            valor_efd,
        ])

    primeira_tabela_fim = ws.max_row
    ws.auto_filter.ref = f"A2:D{primeira_tabela_fim}"
    ws.freeze_panes = "A3"

    ws.append([])
    ws.append([])

    # BLOCO 2 Titulo amarelo
    # título amarelo
    titulo2_row = ws.max_row + 1
    ws.merge_cells(
        start_row=titulo2_row,
        start_column=1,
        end_row=titulo2_row,
        end_column=4,
    )
    ws.cell(titulo2_row, 1).value = (
        "Mismatches detectados (categoria ECD com classificação vs naturezas EFD esperadas)"
    )
    ws.cell(titulo2_row, 1).font = Font(bold=True, color="000000")
    ws.cell(titulo2_row, 1).alignment = Alignment(horizontal="center")
    ws.cell(titulo2_row, 1).fill = PatternFill("solid", fgColor="FFF2CC")

    # cabeçalho cinza
    headers2 = [
        "Categoria ECD",
        "Despesa elegível ECD",
        "Naturezas EFD esperadas",
        "Base declarada nessas naturezas",
    ]

    ws.append(headers2)
    headers2_row = ws.max_row

    for cell in ws[headers2_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")

    for item in iter_items(ctx.get("por_natureza")):
        categoria = item.get("categoria")
        if not categoria:
            continue

        nats = [
            str(nat).zfill(2)
            for nat in (item.get("naturezas_esperadas") or [])
        ]

        base_declarada = sum(
            efd_por_nat.get(nat, 0)
            for nat in nats
        )

        ws.append([
            categoria,
            item.get("valor_ecd"),
            ", ".join(nats),
            base_declarada,
        ])

    ws.append([])
    ws.append([
        "Linhas destacam categorias classificadas na ECD e a cobertura encontrada nas naturezas de crédito esperadas da EFD-Contribuições."
    ])

    for col in ["D", "B"]:
        for row in range(3, ws.max_row + 1):
            ws[f"{col}{row}"].number_format = '#,##0.00'

    autosize_columns(ws)