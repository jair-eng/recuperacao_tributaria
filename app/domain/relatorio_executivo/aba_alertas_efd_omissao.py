from __future__ import annotations

from openpyxl import Workbook

from app.utils.ecd_gap_utils import (
    classificar_alerta_natureza,
    descricao_alerta_natureza,
    montar_base_por_natureza_trimestral,
)
from app.utils.excel import criar_aba_generica, autosize_columns


def criar_aba_alertas_efd_omissao(
    wb: Workbook,
    ctx: dict,
) -> None:
    rows = []

    mapa_nat = ctx.get("mapa_nat_bc_cred") or {}
    agregado_nat = montar_base_por_natureza_trimestral(ctx)

    for item in sorted(
        agregado_nat.values(),
        key=lambda x: (
            x.get("trimestre") or "",
            x.get("nat_bc_cred") or "",
        ),
    ):
        tipo_alerta = classificar_alerta_natureza(item)

        if not tipo_alerta:
            continue

        nat = str(item.get("nat_bc_cred") or "00").zfill(2)

        rows.append(
            {
                "Ano-Trimestre": item["trimestre"],
                "Tipo Alerta": tipo_alerta,
                "Cód. Nat.": nat,
                "Natureza": mapa_nat.get(nat, ""),
                "Valor ECD": item["valor_ecd"],
                "Total Documentado": item["total_documentado"],
                "Bloco M": item["valor_bloco_m"],
                "GAP ECD": item["gap_ecd"],
                "Cobertura Documental %": item["cobertura_documental_pct"],
                "Status": item["status"],
                "Descrição": descricao_alerta_natureza(tipo_alerta),
            }
        )

    ctx["alertas_efd_omissao"] = rows

    nome_aba = "Alertas Efd Omissao"

    criar_aba_generica(
        wb,
        nome_aba=nome_aba,
        headers=[
            "Ano-Trimestre",
            "Tipo Alerta",
            "Cód. Nat.",
            "Natureza",
            "Valor ECD",
            "Total Documentado",
            "Bloco M",
            "GAP ECD",
            "Cobertura Documental %",
            "Status",
            "Descrição",
        ],
        rows=rows,
    )

    ws = wb[nome_aba]

    # Congela primeira linha e três primeiras colunas
    ws.freeze_panes = "D2"

    # Formata valores
    for row in range(2, ws.max_row + 1):
        for col in ["E", "F", "G", "H"]:
            ws[f"{col}{row}"].number_format = '#,##0.00'

        # Cobertura Documental %
        ws[f"I{row}"].number_format = '0.00%'

    ws.auto_filter.ref = f"A1:K{ws.max_row}"

    autosize_columns(ws)